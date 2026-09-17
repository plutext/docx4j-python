"""Parity with real python-docx: python-docx is the judge of what this engine wrote.

CR-003 section 3.6, Phase J, section 20. The subset test
(``test_python_docx_subset.py``) asserts the facade has python-docx's members;
this one asserts the facade *does the same thing*. One script is run twice ---
once through :mod:`docx4j_py.docx`, once through python-docx itself --- and the
two files are then read back **by python-docx**, so the comparison is
python-docx's reading of both and not this engine's reading of its own output.

python-docx is not a dependency of this package (``pyproject.toml``'s ``parity``
extra names it), so this module skips in ``.venv-fork`` and in CI. Run it from a
throwaway virtual environment that has python-docx **and** this package:

    uv venv /tmp/pydocx
    VIRTUAL_ENV=/tmp/pydocx uv pip install python-docx==1.2.0 pytest \\
        -e ~/git/docx4j-xsdata -e .
    /tmp/pydocx/bin/python -m pytest tests/content/test_python_docx_parity.py

The result of the run of 2026-09-18 is in CR-003 section 20.
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path
from typing import Any

import pytest

pytest.importorskip(
    "docx",
    reason="python-docx is not installed here, and is not a dependency of this package",
)

import docx
import docx.enum.section
import docx.enum.table
import docx.enum.text
import docx.shared

import docx4j_py.docx as facade

pytestmark = pytest.mark.skipif(
    docx.__version__ != "1.2.0",
    reason=f"the parity script is written against python-docx 1.2.0, not {docx.__version__}",
)


class Api:
    """The handful of names the script uses, from one library or the other."""

    def __init__(self, name: str, **members: Any) -> None:
        """Name the library and carry its names."""
        self.name = name
        for key, value in members.items():
            setattr(self, key, value)

    def __repr__(self) -> str:
        """``<Api docx4j_py.docx>``."""
        return f"<Api {self.name}>"


DOCX4J = Api(
    "docx4j_py.docx",
    Document=facade.Document,
    Pt=facade.Pt,
    Inches=facade.Inches,
    RGBColor=facade.RGBColor,
    WD_ALIGN_PARAGRAPH=facade.WD_ALIGN_PARAGRAPH,
    WD_ORIENT=facade.WD_ORIENT,
    WD_SECTION=facade.WD_SECTION,
    WD_TABLE_ALIGNMENT=facade.WD_TABLE_ALIGNMENT,
    WD_TAB_ALIGNMENT=facade.WD_TAB_ALIGNMENT,
    WD_TAB_LEADER=facade.WD_TAB_LEADER,
)

PYTHON_DOCX = Api(
    "python-docx",
    Document=docx.Document,
    Pt=docx.shared.Pt,
    Inches=docx.shared.Inches,
    RGBColor=docx.shared.RGBColor,
    WD_ALIGN_PARAGRAPH=docx.enum.text.WD_ALIGN_PARAGRAPH,
    WD_ORIENT=docx.enum.section.WD_ORIENT,
    WD_SECTION=docx.enum.section.WD_SECTION,
    WD_TABLE_ALIGNMENT=docx.enum.table.WD_TABLE_ALIGNMENT,
    WD_TAB_ALIGNMENT=docx.enum.text.WD_TAB_ALIGNMENT,
    WD_TAB_LEADER=docx.enum.text.WD_TAB_LEADER,
)

#: The table the script builds, as a list of dicts, which is how tutorials do it.
ROWS = [
    {"Region": "North", "Quarter": "Q1", "Revenue": "120"},
    {"Region": "South", "Quarter": "Q1", "Revenue": "98"},
]

FIXED = dt.datetime(2026, 9, 18, 10, 0, tzinfo=dt.UTC)


def build(api: Api, target: Path) -> None:
    """One script, written once, run against either library.

    Deliberately ordinary python-docx: nothing here is aware of which library
    it is talking to, which is the whole point of the comparison.
    """
    document = api.Document()

    document.core_properties.title = "Quarterly report"
    document.core_properties.author = "Ada Lovelace"
    document.core_properties.subject = "Numbers"
    document.core_properties.keywords = "quarter, revenue"
    document.core_properties.created = FIXED
    document.core_properties.modified = FIXED

    document.add_heading("Quarterly report", level=0)
    document.add_heading("Summary", level=1)

    paragraph = document.add_paragraph("Revenue rose in ")
    strong = paragraph.add_run("every region")
    strong.bold = True
    paragraph.add_run(" except ")
    weak = paragraph.add_run("the lazy one")
    weak.italic = True
    weak.font.size = api.Pt(14)
    weak.font.name = "Consolas"
    weak.font.color.rgb = api.RGBColor(0xC0, 0x00, 0x00)
    paragraph.add_run(".")
    paragraph.alignment = api.WD_ALIGN_PARAGRAPH.JUSTIFY

    document.add_heading("By region", level=2)
    headers = list(ROWS[0])
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = api.WD_TABLE_ALIGNMENT.CENTER
    for index, name in enumerate(headers):
        cell = table.cell(0, index)
        cell.text = name
        cell.paragraphs[0].runs[0].bold = True
    for row in ROWS:
        cells = table.add_row().cells
        for index, name in enumerate(headers):
            cells[index].text = row[name]

    ledger = document.add_paragraph("Item\tPrice")
    ledger.paragraph_format.tab_stops.add_tab_stop(api.Inches(4), api.WD_TAB_ALIGNMENT.RIGHT, api.WD_TAB_LEADER.DOTS)
    ledger.paragraph_format.tab_stops.add_tab_stop(api.Inches(1))

    document.add_page_break()
    document.add_paragraph("A bulleted point", style="List Bullet")
    document.add_paragraph("A numbered step", style="List Number")

    section = document.add_section(api.WD_SECTION.NEW_PAGE)
    section.orientation = api.WD_ORIENT.LANDSCAPE
    section.page_width = api.Inches(11)
    section.page_height = api.Inches(8.5)
    section.left_margin = api.Inches(1)
    document.add_paragraph("Landscape at the end.")

    document.save(str(target))


@pytest.fixture(scope="module")
def built(tmp_path_factory) -> dict[str, Any]:
    """Both files, opened by **python-docx**."""
    directory = tmp_path_factory.mktemp("parity")
    ours, theirs = directory / "docx4j.docx", directory / "pydocx.docx"
    build(DOCX4J, ours)
    build(PYTHON_DOCX, theirs)
    return {
        "ours": docx.Document(str(ours)),
        "theirs": docx.Document(str(theirs)),
        "paths": (ours, theirs),
    }


def test_both_files_open_in_python_docx(built):
    """The first thing a parity test proves: python-docx can read what we wrote."""
    assert built["ours"] is not None
    for path in built["paths"]:
        assert path.stat().st_size > 2000


def test_the_paragraph_texts_are_the_same(built):
    ours = [p.text for p in built["ours"].paragraphs]
    theirs = [p.text for p in built["theirs"].paragraphs]
    assert ours == theirs


def test_the_paragraph_styles_are_the_same(built):
    ours = [p.style.name for p in built["ours"].paragraphs]
    theirs = [p.style.name for p in built["theirs"].paragraphs]
    assert ours == theirs


def test_the_paragraph_alignments_are_the_same(built):
    ours = [p.alignment for p in built["ours"].paragraphs]
    theirs = [p.alignment for p in built["theirs"].paragraphs]
    assert ours == theirs


def test_the_run_formatting_is_the_same(built):
    def formatting(document: Any) -> list[Any]:
        return [
            (run.text, run.bold, run.italic, run.font.size, run.font.name, run.font.color.rgb)
            for paragraph in document.paragraphs
            for run in paragraph.runs
        ]

    assert formatting(built["ours"]) == formatting(built["theirs"])


def test_the_table_reads_back_the_same(built):
    def shape(document: Any) -> Any:
        table = document.tables[0]
        return (
            len(table.rows),
            len(table.columns),
            [[cell.text for cell in row.cells] for row in table.rows],
            table.style.name,
            table.alignment,
            [row.cells[0].paragraphs[0].runs[0].bold for row in table.rows[:1]],
        )

    assert shape(built["ours"]) == shape(built["theirs"])


def test_the_sections_start_type_and_orientation_are_the_same(built):
    """Every section's break kind and orientation, which the script sets."""

    def sections(document: Any) -> list[Any]:
        return [(s.start_type, s.orientation) for s in document.sections]

    assert sections(built["ours"]) == sections(built["theirs"])


def test_the_page_setup_the_script_wrote_is_the_same_to_the_emu(built):
    """The **last** section, whose page size and margin the script set itself.

    The first section's page setup comes from each library's default template
    and deliberately differs (departure 1 of CR-003 section 20); the test below
    shows that difference on purpose.
    """

    def setup(document: Any) -> Any:
        section = document.sections[-1]
        return (
            section.orientation,
            section.page_width,
            section.page_height,
            section.left_margin,
        )

    assert setup(built["ours"]) == setup(built["theirs"])


def test_the_core_properties_read_back_the_same(built):
    def properties(document: Any) -> Any:
        core = document.core_properties
        return (
            core.title,
            core.author,
            core.subject,
            core.keywords,
            core.created,
            core.modified,
        )

    assert properties(built["ours"]) == properties(built["theirs"])


def test_the_page_break_is_where_python_docx_puts_one(built):
    def breaks(document: Any) -> list[int]:
        return [
            index
            for index, paragraph in enumerate(document.paragraphs)
            if paragraph._p.xpath('.//w:br[@w:type="page"]')
        ]

    assert breaks(built["ours"]) == breaks(built["theirs"])


def test_the_inner_content_order_is_the_same(built):
    def kinds(document: Any) -> list[str]:
        return [type(item).__name__ for item in document.iter_inner_content()]

    assert kinds(built["ours"]) == kinds(built["theirs"])


def test_what_the_two_default_templates_differ_in_is_only_the_template(built):
    """The one recorded departure, shown rather than asserted away.

    ``Document()`` here is ``create_package()`` --- A4, 2.54 cm margins,
    docx4j's default styles --- and python-docx's is its own US Letter
    template. The script above sets the page size of the section it adds, so
    the *last* section matches; the first does not, and that is departure 1 of
    CR-003 section 20.
    """
    ours, theirs = built["ours"].sections[0], built["theirs"].sections[0]
    assert ours.page_width != theirs.page_width
    assert round(ours.page_width.cm, 1) == 21.0, "A4"
    assert round(theirs.page_width.inches, 1) == 8.5, "US Letter"
    assert built["ours"].sections[-1].page_width == built["theirs"].sections[-1].page_width


def test_the_tab_stops_read_back_the_same(built):
    """Position, alignment and leader of each stop, as python-docx reads them."""

    def stops(document):
        for paragraph in document.paragraphs:
            if paragraph.text.startswith("Item\t"):
                return [(s.position, s.alignment, s.leader) for s in paragraph.paragraph_format.tab_stops]
        raise AssertionError("the ledger paragraph is missing")

    assert stops(built["ours"]) == stops(built["theirs"])
    assert len(stops(built["ours"])) == 2
